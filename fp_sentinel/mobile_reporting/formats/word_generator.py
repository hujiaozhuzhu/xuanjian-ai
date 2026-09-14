"""Word 安全审计报告生成器（玄鉴AI / fp_sentinel）。

基于 python-docx 生成参照绿盟/启明星辰商业渗透测试报告版式的 ``.docx``
报告，包含封面、目录域、概述与声明、执行摘要（严重度分布着色统计表）、
逐漏洞详细分析（信息表 / 证据 / 复现步骤 / 截图 / POC / EXP / 修复建议 /
参考链接）、测试环境与术语表附录以及文档信息页。

数据来源兼容两种模型包：

- 优先使用 ``fp_sentinel.mobile_reporting.models.report_models``；
- 未就绪时回退到 :mod:`fp_sentinel.mobile_reporting.formats._fallback_models`，
  字段读取统一通过 :func:`get_attr` 鸭子类型访问。

输出路径受 S7 红线约束：仅允许写入白名单根目录内的 ``.docx`` 文件。
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:  # python-docx 不可用时给出明确指引
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - 环境依赖
    raise ImportError(
        "WordGenerator 依赖 python-docx，请先安装: pip install python-docx"
    ) from exc

from .base_generator import BaseReportGenerator
from ._fallback_models import get_attr

try:  # 正式模型由 models 子包提供；未就绪时使用回退定义
    from ..models.report_models import MobileSecurityReport
except ImportError:  # pragma: no cover - 依赖并行开发的模型包
    from ._fallback_models import MobileSecurityReport

__all__ = ["WordGenerator", "ReportGenerationError"]

logger = logging.getLogger(__name__)

# ─────────────────────────────── 常量 ───────────────────────────────

#: 严重度从高到低的固定展示顺序
_SEVERITY_ORDER: Tuple[str, ...] = (
    "CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO",
)

#: 严重度单元格底色（十六进制 RGB，不含 #）
_SEVERITY_BG: Dict[str, str] = {
    "CRITICAL": "8B1A1A",  # 深红
    "HIGH": "C00000",  # 红
    "MEDIUM": "E36C0A",  # 橙
    "LOW": "FFC000",  # 黄
    "INFO": "2E74B5",  # 蓝
}

#: 严重度单元格字色（与底色对比调整）
_SEVERITY_FG: Dict[str, RGBColor] = {
    "CRITICAL": RGBColor(0xFF, 0xFF, 0xFF),
    "HIGH": RGBColor(0xFF, 0xFF, 0xFF),
    "MEDIUM": RGBColor(0xFF, 0xFF, 0xFF),
    "LOW": RGBColor(0x00, 0x00, 0x00),
    "INFO": RGBColor(0xFF, 0xFF, 0xFF),
}

_HEADER_BG = "D9D9D9"  # 表头底色（浅灰）
_CODE_BG = "F2F2F2"  # 代码块底色（浅灰）
_EAST_ASIA_BODY = "宋体"  # 中文正文
_EAST_ASIA_HEAD = "黑体"  # 中文标题
_ASCII_BODY = "Times New Roman"
_ASCII_HEAD = "Arial"
_ASCII_MONO = "Consolas"
_POC_DANGER_MAX_LINES = 10  # CRITICAL 级 POC 展示行数上限

_DISCLAIMER = (
    "本报告由玄鉴AI（fp_sentinel）自动化安全评估平台生成，仅用于授权范围内的"
    "安全测试与风险评估。报告结论基于测试时点的应用版本与测试环境，不代表该"
    "应用在其他版本、渠道或运行环境下的安全状况。未经编制方书面授权，任何人"
    "不得复制、传播本报告的全部或部分内容，不得利用报告中披露的漏洞信息对目"
    "标系统实施任何未授权的攻击行为。因违反上述约定使用本报告所引发的一切后"
    "果，由行为人自行承担，编制方不承担相关责任。"
)

_GLOSSARY: List[Tuple[str, str]] = [
    ("APK", "Android Package，安卓应用安装包格式。"),
    ("CWE", "Common Weakness Enumeration，通用缺陷枚举编号体系。"),
    ("MASVS", "OWASP 移动应用安全验证标准（v2.0）。"),
    ("POC", "Proof of Concept，无害化概念验证代码。"),
    ("EXP", "Exploit，完整利用代码，仅限授权测试环境使用。"),
    ("Frida", "跨平台动态插桩框架，用于运行时 Hook 与行为观测。"),
    ("静态分析", "不运行目标程序，通过反编译与规则/语义扫描发现缺陷。"),
    ("动态调试", "在运行态观测与篡改程序行为，验证漏洞真实可达性。"),
    ("复现", "按既定步骤重新触发漏洞并固定证据的过程。"),
    ("密级", "报告的保密等级标注，控制报告的分发范围。"),
]


class ReportGenerationError(RuntimeError):
    """Word 报告生成或生成后自检失败。"""


# ─────────────────────────── XML / 样式工具 ───────────────────────────


def _set_style_east_asia(style: Any, east_asia: str) -> None:
    """为样式补充中文（eastAsia）字体设置。"""
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def _set_run_east_asia(run: Any, east_asia: str) -> None:
    """为单个文本运行补充中文（eastAsia）字体设置。"""
    rpr = run._element.get_or_add_rPr()  # noqa: SLF001 - python-docx 内部接口
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def _shade_paragraph(paragraph: Any, fill: str) -> None:
    """为段落追加底纹（w:shd）。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)  # noqa: SLF001


def _shade_cell(cell: Any, fill: str) -> None:
    """为表格单元格追加底色（w:shd）。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)  # noqa: SLF001


def _add_field_run(paragraph: Any, instruction: str) -> None:
    """在段落中插入 Word 域（如 PAGE / NUMPAGES 页码域）。"""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._element.append(begin)  # noqa: SLF001
    run._element.append(instr)  # noqa: SLF001
    run._element.append(end)  # noqa: SLF001


def _write_cell(
    cell: Any,
    text: str,
    bold: bool = False,
    color: Optional[RGBColor] = None,
    mono: bool = False,
) -> None:
    """向单元格写入文本并设置字体样式。"""
    run = cell.paragraphs[0].add_run(text)
    if mono:
        run.font.name = _ASCII_MONO
        _set_run_east_asia(run, _EAST_ASIA_BODY)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_table(doc: Any, headers: List[str], rows: List[List[str]]) -> Any:
    """创建统一边框、表头带底色的表格。"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        _write_cell(cell, header, bold=True)
        _shade_cell(cell, _HEADER_BG)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            _write_cell(cells[idx], str(value))
    return table


def _colorize_severity_cell(
    table: Any, row_idx: int, col_idx: int, severity: str
) -> None:
    """将指定单元格按严重度着色（底色 + 字色）。"""
    cell = table.rows[row_idx].cells[col_idx]
    _shade_cell(cell, _SEVERITY_BG.get(severity, "808080"))
    fg = _SEVERITY_FG.get(severity, RGBColor(0x00, 0x00, 0x00))
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = fg
            run.bold = True


def _add_code_block(doc: Any, code: str, size: float = 9.0) -> None:
    """添加等宽字体、浅灰底纹的代码块段落。"""
    lines = code.splitlines() or [""]
    for line in lines:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line if line.strip() else " ")
        run.font.name = _ASCII_MONO
        _set_run_east_asia(run, _EAST_ASIA_BODY)
        run.font.size = Pt(size)
        _shade_paragraph(paragraph, _CODE_BG)
        pf = paragraph.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0


def _add_caption(doc: Any, text: str) -> None:
    """添加居中图注（图 X-Y 标题）。"""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(9)
    run.bold = True
    _set_run_east_asia(run, _EAST_ASIA_BODY)


def _add_placeholder(doc: Any, text: str) -> None:
    """添加占位/警示段落（如截图缺失说明）。"""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    _set_run_east_asia(run, _EAST_ASIA_BODY)


# ───────────────────────────── 生成器主体 ─────────────────────────────


class WordGenerator(BaseReportGenerator):
    """基于 python-docx 的 Word 渗透测试报告生成器。

    Args:
        allowed_roots: 输出路径白名单根目录集合（S7 红线）。
        allowed_suffixes: 允许的输出后缀；默认仅 ``.docx``。
    """

    DEFAULT_ALLOWED_SUFFIXES = frozenset({".docx"})

    def __init__(
        self,
        allowed_roots: Optional[List[Any]] = None,
        allowed_suffixes: Optional[List[str]] = None,
    ) -> None:
        super().__init__(
            allowed_roots=allowed_roots, allowed_suffixes=allowed_suffixes
        )

    # ────────────────────────── 对外接口 ──────────────────────────

    def get_format_name(self) -> str:
        """返回格式名称。"""
        return "Word (.docx)"

    def generate(
        self, report: MobileSecurityReport, output_path: Path
    ) -> Path:
        """生成 Word 报告并执行生成后自检。

        Args:
            report: 移动安全评估报告数据（正式模型或回退模型均可）。
            output_path: 输出 ``.docx`` 路径（必须位于白名单根目录内）。

        Returns:
            实际写入的产物绝对路径。

        Raises:
            PathNotAllowedError: 输出路径违反 S7 白名单红线。
            ReportGenerationError: 文档构建、写入或自检失败。
        """
        target = self.validate_output_path(Path(output_path))
        logger.info("开始生成 Word 报告: %s", target)
        doc = self._build_document(report)
        try:
            doc.save(str(target))
        except OSError as exc:
            raise ReportGenerationError(
                f"写入 Word 文件失败: {target} ({exc})"
            ) from exc
        self._self_check(target)
        logger.info("Word 报告生成完成: %s", target)
        return target

    # ────────────────────────── 文档构建 ──────────────────────────

    def _build_document(self, report: MobileSecurityReport) -> Any:
        """构建完整文档对象。"""
        try:
            doc = Document()
            self._configure_page(doc)
            self._configure_styles(doc)
            self._configure_header_footer(doc, report)
            self._build_cover(doc, report)
            self._build_toc(doc)
            self._build_overview(doc, report)
            self._build_summary(doc, report)
            self._build_findings(doc, report)
            self._build_appendix_environment(doc, report)
            self._build_appendix_glossary(doc)
            self._build_document_info(doc, report)
            return doc
        except ReportGenerationError:
            raise
        except Exception as exc:
            raise ReportGenerationError(f"构建 Word 文档失败: {exc}") from exc

    def _configure_page(self, doc: Any) -> None:
        """配置 A4 页面与页边距。"""
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    def _configure_styles(self, doc: Any) -> None:
        """配置正文（宋体）与各级标题（黑体）样式。"""
        normal = doc.styles["Normal"]
        normal.font.name = _ASCII_BODY
        normal.font.size = Pt(10.5)
        _set_style_east_asia(normal, _EAST_ASIA_BODY)

        heading_sizes = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 12}
        for name, size in heading_sizes.items():
            style = doc.styles[name]
            style.font.name = _ASCII_HEAD
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            _set_style_east_asia(style, _EAST_ASIA_HEAD)

    def _configure_header_footer(
        self, doc: Any, report: MobileSecurityReport
    ) -> None:
        """页眉显示报告标题，页脚插入页码域。"""
        title = str(get_attr(report, "title", "移动应用渗透测试报告") or "移动应用渗透测试报告")
        section = doc.sections[0]

        header_para = section.header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run(title)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        _set_run_east_asia(run, _EAST_ASIA_BODY)

        footer_para = section.footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.add_run("第 ")
        _add_field_run(footer_para, "PAGE")
        footer_para.add_run(" 页 / 共 ")
        _add_field_run(footer_para, "NUMPAGES")
        footer_para.add_run(" 页")

    # ────────────────────────── 封面与目录 ──────────────────────────

    def _build_cover(self, doc: Any, report: MobileSecurityReport) -> None:
        """构建封面页：标题、密级、目标 APP 信息表、日期、版本、编制人。"""
        classification = str(
            get_attr(report, "classification", "内部资料") or "内部资料"
        )
        title = str(get_attr(report, "title", "移动应用渗透测试报告") or "移动应用渗透测试报告")
        subtitle = str(get_attr(report, "subtitle", "") or "")
        company = str(get_attr(report, "company", "玄鉴信息安全技术有限公司") or "")
        version = str(get_attr(report, "report_version", "V1.0") or "V1.0")
        analyst = str(get_attr(report, "analyst", "") or "玄鉴AI安全评估组")
        scan_date = str(get_attr(report, "scan_date", "") or "")
        start_time = str(get_attr(report, "test_start_time", "") or "")
        end_time = str(get_attr(report, "test_end_time", "") or "")

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"密级：{classification}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        _set_run_east_asia(run, _EAST_ASIA_HEAD)

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(title)
        run.font.name = _ASCII_HEAD
        _set_run_east_asia(run, _EAST_ASIA_HEAD)
        run.font.size = Pt(26)
        run.bold = True

        if subtitle:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(subtitle)
            run.font.size = Pt(12)
            _set_run_east_asia(run, _EAST_ASIA_BODY)

        doc.add_paragraph()
        _add_table(doc, ["项目", "内容"], self._cover_target_rows(report))

        for label, value in (
            ("测试日期", scan_date or f"{start_time} 至 {end_time}"),
            ("报告版本", version),
            ("编制人", analyst),
            ("编制单位", company),
        ):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(f"{label}：{value}")
            run.font.size = Pt(12)
            _set_run_east_asia(run, _EAST_ASIA_BODY)

        doc.add_page_break()

    def _cover_target_rows(
        self, report: MobileSecurityReport
    ) -> List[List[str]]:
        """构造封面目标 APP 信息表数据行。"""
        target = get_attr(report, "target", None)
        if target is None:
            return [("目标应用", "未提供")]
        rows = [
            ("应用名称", str(get_attr(target, "app_name", "") or "N/A")),
            ("包名", str(get_attr(target, "package_name", "") or "N/A")),
            ("版本名称", str(get_attr(target, "version_name", "") or "N/A")),
            (
                "版本号(versionCode)",
                str(get_attr(target, "version_code", "") or "N/A"),
            ),
            (
                "minSdk / targetSdk",
                "{} / {}".format(
                    get_attr(target, "min_sdk", "") or "N/A",
                    get_attr(target, "target_sdk", "") or "N/A",
                ),
            ),
            ("文件大小", str(get_attr(target, "file_size", "") or "N/A")),
        ]
        sha256 = str(get_attr(target, "sha256", "") or "")
        if sha256:
            rows.append(("SHA256", sha256))
        return rows

    def _build_toc(self, doc: Any) -> None:
        """构建目录页：插入 TOC 域与更新域提示。"""
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run("目  录")
        run.font.name = _ASCII_HEAD
        _set_run_east_asia(run, _EAST_ASIA_HEAD)
        run.font.size = Pt(16)
        run.bold = True

        toc_para = doc.add_paragraph()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
        inner_run = OxmlElement("w:r")
        inner_text = OxmlElement("w:t")
        inner_text.text = "（目录域：请在 Word 中按 F9 更新域以生成目录）"
        inner_run.append(inner_text)
        fld.append(inner_run)
        toc_para._p.append(fld)  # noqa: SLF001

        hint = doc.add_paragraph()
        hint.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hint.add_run("提示：本目录为 Word 域，打开文档后按 F9（或右键“更新域”）即可刷新。")
        run.font.size = Pt(9)
        run.italic = True
        _set_run_east_asia(run, _EAST_ASIA_BODY)
        doc.add_page_break()

    # ────────────────────────── 章节内容 ──────────────────────────

    def _build_overview(self, doc: Any, report: MobileSecurityReport) -> None:
        """第 1 章：概述与声明（测试范围 / 测试方法 / 免责声明）。"""
        doc.add_heading("1. 概述与声明", level=1)

        background = str(get_attr(report, "project_background", "") or "")
        if background:
            doc.add_heading("1.1 项目背景", level=2)
            doc.add_paragraph(background)

        doc.add_heading("1.2 测试范围", level=2)
        scope = get_attr(report, "test_scope", []) or []
        if scope:
            for item in scope:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            _add_placeholder(doc, "未提供测试范围信息。")

        doc.add_heading("1.3 测试方法", level=2)
        methods = get_attr(report, "test_methods", []) or []
        if methods:
            for item in methods:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            _add_placeholder(doc, "未提供测试方法说明。")

        doc.add_heading("1.4 免责声明", level=2)
        doc.add_paragraph(_DISCLAIMER)

    def _build_summary(self, doc: Any, report: MobileSecurityReport) -> None:
        """第 2 章：执行摘要（总体评级 / 严重度分布 / 关键发现 Top）。"""
        doc.add_heading("2. 执行摘要", level=1)

        findings = self._get_findings(report)
        counts = self._severity_counts(findings)

        doc.add_heading("2.1 总体风险评级结论", level=2)
        doc.add_paragraph(self._overall_rating(counts))
        assessment = str(get_attr(report, "overall_assessment", "") or "")
        if assessment:
            doc.add_paragraph(assessment)

        doc.add_heading("2.2 严重度分布统计", level=2)
        rows: List[List[str]] = []
        for severity in _SEVERITY_ORDER:
            rows.append([severity, str(counts.get(severity, 0))])
        rows.append(["合计", str(len(findings))])
        table = _add_table(doc, ["严重程度", "数量"], rows)
        for row_idx, severity in enumerate(_SEVERITY_ORDER, start=1):
            _colorize_severity_cell(table, row_idx, 0, severity)

        doc.add_heading("2.3 关键发现 Top 列表", level=2)
        top_findings = self._sorted_findings(findings)[:5]
        if top_findings:
            for idx, vuln in enumerate(top_findings, start=1):
                severity = str(
                    get_attr(vuln, "severity", "INFO") or "INFO"
                ).upper()
                vuln_id = str(get_attr(vuln, "vuln_id", "") or "N/A")
                title = str(get_attr(vuln, "title", "") or "未命名漏洞")
                doc.add_paragraph(
                    f"{idx}. [{severity}] {vuln_id} {title}",
                    style="List Number",
                )
        else:
            _add_placeholder(doc, "本次测试未发现安全问题。")

    def _build_findings(self, doc: Any, report: MobileSecurityReport) -> None:
        """第 3 章：详细漏洞分析，每个 finding 独立小节。"""
        doc.add_heading("3. 详细漏洞分析", level=1)
        findings = self._get_findings(report)
        if not findings:
            _add_placeholder(doc, "本次测试未发现需要详细分析的安全问题。")
            return
        for index, vuln in enumerate(findings, start=1):
            self._build_finding_section(doc, vuln, index)

    def _build_finding_section(self, doc: Any, vuln: Any, index: int) -> None:
        """构建单个漏洞的详细分析小节。"""
        vuln_id = str(get_attr(vuln, "vuln_id", "") or f"FINDING-{index:03d}")
        title = str(get_attr(vuln, "title", "") or "未命名漏洞")
        severity = str(get_attr(vuln, "severity", "INFO") or "INFO").upper()

        doc.add_heading(f"3.{index} {vuln_id} {title}", level=2)

        doc.add_heading(f"3.{index}.1 漏洞基本信息", level=3)
        info_rows = [
            ("风险等级", severity),
            ("CWE 编号", str(get_attr(vuln, "cwe", "") or "N/A")),
            ("OWASP MASVS", str(get_attr(vuln, "masvs", "") or "N/A")),
            ("置信度", str(get_attr(vuln, "confidence", "Firm") or "Firm")),
            ("漏洞类别", str(get_attr(vuln, "category", "") or "N/A")),
        ]
        table = _add_table(doc, ["项目", "内容"], info_rows)
        _colorize_severity_cell(table, 1, 0, severity)

        components = get_attr(vuln, "components", []) or []
        if components:
            doc.add_paragraph("涉及组件：" + "、".join(str(c) for c in components))

        doc.add_heading(f"3.{index}.2 漏洞描述", level=3)
        doc.add_paragraph(str(get_attr(vuln, "description", "") or "无描述。"))
        impact = str(get_attr(vuln, "impact", "") or "")
        if impact:
            doc.add_paragraph(f"危害影响：{impact}")

        self._build_evidence(doc, vuln, index)
        self._build_reproduction(doc, vuln, index)
        self._build_screenshots(doc, vuln, index)
        self._build_poc(doc, vuln, index, severity)
        self._build_exp(doc, vuln, index)
        self._build_fix_and_refs(doc, vuln, index)

    def _build_evidence(self, doc: Any, vuln: Any, index: int) -> None:
        """漏洞证据：代码位置 + 等宽字体、浅灰底纹代码片段。"""
        doc.add_heading(f"3.{index}.3 证据列表", level=3)
        evidence = get_attr(vuln, "evidence", []) or []
        if not evidence:
            _add_placeholder(doc, "无固定化证据。")
            return
        for seq, item in enumerate(evidence, start=1):
            location = str(get_attr(item, "location", "") or "未知位置")
            content = str(get_attr(item, "content", "") or "")
            description = str(get_attr(item, "description", "") or "")
            doc.add_paragraph(f"证据 {seq} 位置：{location}")
            if content:
                _add_code_block(doc, content)
            if description:
                doc.add_paragraph(f"说明：{description}")

    def _build_reproduction(self, doc: Any, vuln: Any, index: int) -> None:
        """复现步骤：操作 → 预期结果 → 实际结果 → 复现命令。"""
        doc.add_heading(f"3.{index}.4 复现步骤", level=3)
        steps = get_attr(vuln, "reproduction_steps", []) or []
        if not steps:
            _add_placeholder(doc, "未提供复现步骤。")
            return
        for seq, step in enumerate(steps, start=1):
            action = str(get_attr(step, "action", "") or "")
            expected = str(get_attr(step, "expected", "") or "")
            actual = str(get_attr(step, "actual", "") or "")
            command = str(get_attr(step, "command", "") or "")
            para = doc.add_paragraph()
            run = para.add_run(f"步骤 {seq}：操作——{action}")
            run.bold = True
            _set_run_east_asia(run, _EAST_ASIA_BODY)
            doc.add_paragraph(f"预期结果：{expected or 'N/A'}")
            doc.add_paragraph(f"实际结果：{actual or 'N/A'}")
            doc.add_paragraph("复现命令：")
            if command:
                _add_code_block(doc, command)

    def _build_screenshots(self, doc: Any, vuln: Any, index: int) -> None:
        """截图：嵌入真实图片并加图注；缺失时插入占位说明。"""
        doc.add_heading(f"3.{index}.5 截图", level=3)
        screenshots = get_attr(vuln, "screenshots", []) or []
        if not screenshots:
            _add_placeholder(doc, "本漏洞无截图证据。")
            return
        for fig_no, shot in enumerate(screenshots, start=1):
            path_text = str(get_attr(shot, "path", "") or "")
            description = str(get_attr(shot, "description", "") or "截图")
            if not path_text:
                _add_placeholder(doc, f"截图缺失：未提供截图路径（{description}）。")
                continue
            image_path = Path(path_text)
            if not image_path.is_file():
                _add_placeholder(
                    doc,
                    f"截图缺失：文件不存在（{image_path}）。"
                    "请补齐截图文件后重新生成报告。",
                )
                continue
            try:
                doc.add_picture(str(image_path), width=Cm(14.0))
            except Exception as exc:
                logger.warning("截图嵌入失败 %s: %s", image_path, exc)
                _add_placeholder(
                    doc,
                    f"截图缺失：图片无法解析或嵌入（{image_path}）：{exc}",
                )
                continue
            _add_caption(doc, f"图 {index}-{fig_no} {description}")

    def _build_poc(
        self, doc: Any, vuln: Any, index: int, severity: str
    ) -> None:
        """POC 脚本：等宽代码块；CRITICAL 级仅展示前 10 行并附警示语。"""
        doc.add_heading(f"3.{index}.6 POC 脚本", level=3)
        poc = str(get_attr(vuln, "poc", "") or "")
        if not poc.strip():
            _add_placeholder(doc, "本次测试未编写 POC 脚本。")
            return
        if severity == "CRITICAL":
            doc.add_paragraph(
                "【警示】该漏洞等级为 CRITICAL，以下仅展示 POC 前 "
                f"{_POC_DANGER_MAX_LINES} 行用于验证说明；完整利用代码不予展示，"
                "仅限授权测试环境内部使用，严禁外传。"
            )
            lines = poc.splitlines()
            truncated = len(lines) > _POC_DANGER_MAX_LINES
            shown = lines[:_POC_DANGER_MAX_LINES]
            _add_code_block(doc, "\n".join(shown))
            if truncated:
                _add_placeholder(
                    doc,
                    f"（内容已截断：完整 POC 共 {len(lines)} 行，"
                    f"仅展示前 {_POC_DANGER_MAX_LINES} 行。）",
                )
            return
        _add_code_block(doc, poc)

    def _build_exp(self, doc: Any, vuln: Any, index: int) -> None:
        """EXP 说明：默认给出受控利用边界说明。"""
        doc.add_heading(f"3.{index}.7 EXP 说明", level=3)
        exp = str(get_attr(vuln, "exp", "") or "")
        if exp.strip():
            doc.add_paragraph(exp)
        else:
            _add_placeholder(
                doc,
                "本次测试未开展深入利用（EXP）。建议在授权范围内结合业务影响"
                "进一步评估实际利用链与危害半径。",
            )

    def _build_fix_and_refs(self, doc: Any, vuln: Any, index: int) -> None:
        """修复建议与参考链接。"""
        doc.add_heading(f"3.{index}.8 修复建议", level=3)
        suggestions = get_attr(vuln, "fix_suggestions", []) or []
        if suggestions:
            for item in suggestions:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            _add_placeholder(doc, "未提供修复建议。")

        doc.add_heading(f"3.{index}.9 参考链接", level=3)
        references = get_attr(vuln, "references", []) or []
        if references:
            for item in references:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            _add_placeholder(doc, "无参考链接。")

    # ────────────────────────── 附录与信息页 ──────────────────────────

    def _build_appendix_environment(
        self, doc: Any, report: MobileSecurityReport
    ) -> None:
        """第 4 章：附录A 测试环境。"""
        doc.add_heading("4. 附录A 测试环境", level=1)
        environment = get_attr(report, "environment", {}) or {}
        if not environment:
            _add_placeholder(doc, "未提供测试环境信息。")
            return
        rows = [[str(k), str(v)] for k, v in environment.items()]
        _add_table(doc, ["环境项", "版本 / 说明"], rows)

    def _build_appendix_glossary(self, doc: Any) -> None:
        """第 5 章：附录B 术语表。"""
        doc.add_heading("5. 附录B 术语表", level=1)
        _add_table(doc, ["术语", "释义"], [[k, v] for k, v in _GLOSSARY])

    def _build_document_info(
        self, doc: Any, report: MobileSecurityReport
    ) -> None:
        """第 6 章：文档信息页（版本记录表）。"""
        doc.add_heading("6. 文档信息", level=1)
        doc.add_heading("6.1 报告基本信息", level=2)
        _add_table(
            doc,
            ["项目", "内容"],
            [
                ("报告标题", str(get_attr(report, "title", "") or "N/A")),
                ("密级", str(get_attr(report, "classification", "") or "N/A")),
                ("编制单位", str(get_attr(report, "company", "") or "N/A")),
                ("报告版本", str(get_attr(report, "report_version", "") or "N/A")),
            ],
        )
        doc.add_heading("6.2 版本记录", level=2)
        history = get_attr(report, "revision_history", None) or []
        if history:
            rows = [
                [
                    str(get_attr(item, "version", "") or ""),
                    str(get_attr(item, "date", "") or ""),
                    str(get_attr(item, "author", "") or ""),
                    str(get_attr(item, "note", "") or ""),
                ]
                for item in history
            ]
        else:
            rows = [
                [
                    str(get_attr(report, "report_version", "") or "V1.0"),
                    str(get_attr(report, "scan_date", "") or "N/A"),
                    str(get_attr(report, "analyst", "") or "玄鉴AI安全评估组"),
                    "首次发布",
                ]
            ]
        _add_table(doc, ["版本", "日期", "编制/修订人", "修订说明"], rows)

    # ────────────────────────── 数据辅助 ──────────────────────────

    @staticmethod
    def _get_findings(report: MobileSecurityReport) -> List[Any]:
        """读取漏洞列表，兼容 vulnerabilities / findings 两种字段名。"""
        findings = get_attr(report, "vulnerabilities", None)
        if not findings:
            findings = get_attr(report, "findings", None)
        return list(findings or [])

    @staticmethod
    def _severity_counts(findings: List[Any]) -> Dict[str, int]:
        """统计各严重度数量。"""
        counts: Dict[str, int] = {name: 0 for name in _SEVERITY_ORDER}
        for item in findings:
            severity = str(
                get_attr(item, "severity", "INFO") or "INFO"
            ).upper()
            counts[severity] = counts.get(severity, 0) + 1
        return counts

    @staticmethod
    def _sorted_findings(findings: List[Any]) -> List[Any]:
        """按严重度从高到低排序。"""
        rank = {name: i for i, name in enumerate(_SEVERITY_ORDER)}
        return sorted(
            findings,
            key=lambda item: rank.get(
                str(get_attr(item, "severity", "INFO") or "INFO").upper(), 99
            ),
        )

    @staticmethod
    def _overall_rating(counts: Dict[str, int]) -> str:
        """根据严重度分布给出总体风险评级结论。"""
        if counts.get("CRITICAL", 0):
            return "高风险：存在严重（CRITICAL）级别漏洞，建议立即启动应急修复。"
        if counts.get("HIGH", 0):
            return "高风险：存在高危漏洞，建议尽快纳入修复计划。"
        if counts.get("MEDIUM", 0):
            return "中风险：存在中危漏洞，建议在下一迭代中修复。"
        if counts.get("LOW", 0):
            return "低风险：仅存在低危问题，建议择期修复。"
        return "低风险：本次测试未发现明显安全问题。"

    # ────────────────────────── 生成后自检 ──────────────────────────

    def _self_check(self, target: Path) -> None:
        """重新打开产物校验段落数与 Heading1 章节数。

        Raises:
            ReportGenerationError: 段落数不超过 50 或 Heading1 数量不为 6。
        """
        try:
            doc = Document(str(target))
        except Exception as exc:
            raise ReportGenerationError(
                f"自检失败：产物 {target} 无法用 python-docx 重新打开 ({exc})"
            ) from exc
        paragraph_count = len(doc.paragraphs)
        h1_texts = [
            p.text for p in doc.paragraphs if p.style.name == "Heading 1"
        ]
        if paragraph_count <= 50 or len(h1_texts) != 6:
            raise ReportGenerationError(
                "自检失败：段落构成不符合预期，"
                f"段落数={paragraph_count}（应>50），"
                f"Heading1数量={len(h1_texts)}（应为6）：{h1_texts}"
            )
        logger.debug(
            "自检通过: 段落数=%d, Heading1=%s", paragraph_count, h1_texts
        )
